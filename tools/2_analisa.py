# -*- coding: utf-8 -*-
"""
2_analisa.py - transaksi.csv -> laporan.html + tracker.xlsx

    python 2_analisa.py transaksi.csv
    python 2_analisa.py transaksi.csv --kategori config\kategori.txt

Hasil:
    laporan.html   satu file, bisa dibuka di HP/laptop, tanpa internet
    tracker.xlsx   Excel: sheet Input + Ringkasan
"""

import argparse
import io
import json
import os
import sys
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import baca_csv, judul, rapikan, rupiah

HERE = os.path.dirname(os.path.abspath(__file__))
AKAR = os.path.dirname(HERE)
KATEGORI_DEFAULT = os.path.join(AKAR, "config", "kategori.txt")


def baca_aturan(path):
    aturan = []
    if not os.path.exists(path):
        return aturan
    with io.open(path, encoding="utf-8") as fh:
        for baris in fh:
            baris = baris.strip()
            if not baris or baris.startswith("#"):
                continue
            bagian = [b.strip() for b in baris.split("|")]
            if len(bagian) < 2:
                continue
            jenis = bagian[2].lower() if len(bagian) > 2 else "belanja"
            aturan.append((bagian[0].lower(), bagian[1], jenis))
    return aturan


def kategori_untuk(teks, aturan, cadangan="Lain-lain"):
    rendah = (teks or "").lower()
    for kunci, kat, jenis in aturan:
        if kunci in rendah:
            return kat, jenis
    return cadangan, "belanja"


def proses(path, path_aturan):
    aturan = baca_aturan(path_aturan)
    mentah = baca_csv(path)

    baris = []
    for r in mentah:
        b = rapikan(r)
        if not b["tanggal"] or b["nominal"] <= 0:
            continue
        kat, jenis = kategori_untuk(b["keterangan"], aturan)
        if b["jenis"].upper() == "MASUK":
            jenis = "masuk"
            if kat == "Lain-lain":
                kat = "Uang Masuk"
        b["kategori"] = kat
        b["jenis_uang"] = jenis
        baris.append(b)

    baris.sort(key=lambda x: (x["tanggal"], x["waktu"] or ""))
    return baris, aturan


def ringkas(baris):
    total = defaultdict(int)
    per_kat = defaultdict(lambda: [0, 0])
    per_bulan = defaultdict(int)
    per_vendor = defaultdict(lambda: [0, 0])

    for b in baris:
        jenis = b["jenis_uang"]
        total[jenis] += b["nominal"]
        if jenis == "belanja":
            per_kat[b["kategori"]][0] += b["nominal"]
            per_kat[b["kategori"]][1] += 1
            per_bulan[b["tanggal"][:7]] += b["nominal"]
            v = b["keterangan"][:32]
            per_vendor[v][0] += b["nominal"]
            per_vendor[v][1] += 1

    return {
        "total": dict(total),
        "per_kat": {k: {"v": v[0], "n": v[1]} for k, v in per_kat.items()},
        "per_bulan": dict(per_bulan),
        "per_vendor": {k: {"v": v[0], "n": v[1]} for k, v in per_vendor.items()},
    }


def tulis_excel(baris, r, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Input"

    kepala = ["Tanggal", "Jam", "Keterangan", "Kategori", "Jenis", "Nominal", "Sumber"]
    ws.append(kepala)
    for c in range(1, len(kepala) + 1):
        sel = ws.cell(row=1, column=c)
        sel.font = Font(bold=True, color="FFFFFF")
        sel.fill = PatternFill("solid", fgColor="1E7B7B")
        sel.alignment = Alignment(horizontal="center")

    for b in baris:
        ws.append([b["tanggal"], b["waktu"], b["keterangan"], b["kategori"],
                   b["jenis_uang"], b["nominal"], b["sumber"]])

    lebar = [12, 7, 46, 24, 10, 14, 14]
    for i, w in enumerate(lebar, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for sel in ws["F"][1:]:
        sel.number_format = '"Rp"#,##0'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:G%d" % (len(baris) + 1)

    rs = wb.create_sheet("Ringkasan")
    rs.append(["Kategori", "Transaksi", "Total"])
    for c in range(1, 4):
        sel = rs.cell(row=1, column=c)
        sel.font = Font(bold=True, color="FFFFFF")
        sel.fill = PatternFill("solid", fgColor="1E7B7B")

    urut = sorted(r["per_kat"].items(), key=lambda x: -x[1]["v"])
    for i, (kat, v) in enumerate(urut, start=2):
        rs.append([kat, v["n"], v["v"]])
        rs.cell(row=i, column=3).number_format = '"Rp"#,##0'
    akhir = len(urut) + 1
    rs.append(["TOTAL", "", "=SUM(C2:C%d)" % akhir])
    rs.cell(row=akhir + 1, column=1).font = Font(bold=True)
    rs.cell(row=akhir + 1, column=3).font = Font(bold=True)
    rs.cell(row=akhir + 1, column=3).number_format = '"Rp"#,##0'
    rs.column_dimensions["A"].width = 28
    rs.column_dimensions["B"].width = 12
    rs.column_dimensions["C"].width = 16

    wb.save(path)
    return path


HTML = u"""<!DOCTYPE html>
<html lang="id"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Laporan Pengeluaran</title>
<style>
:root{--dark:#123b3b;--mid:#1e7b7b;--line:#e3ecec;--muted:#6b8484;--warn:#b45309;--warnbg:#fdf3e3}
*{box-sizing:border-box}
body{margin:0;font-family:system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;
background:#f5f8f8;color:#173}
header{background:linear-gradient(135deg,#123b3b,#1e7b7b);color:#fff;padding:26px 0}
.wrap{max-width:1080px;margin:0 auto;padding:0 18px}
h1{margin:0;font-size:24px;letter-spacing:-.4px}
.sub{margin:6px 0 0;opacity:.85;font-size:13px}
main{padding:20px 0 40px}
.card{background:#fff;border:1px solid var(--line);border-radius:12px;padding:18px;margin-bottom:16px}
.card h2{margin:0 0 4px;font-size:15px;text-transform:uppercase;letter-spacing:.6px;color:var(--muted)}
.kpi{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin-bottom:16px}
.kpi .k{background:#fff;border:1px solid var(--line);border-left:4px solid var(--mid);
border-radius:10px;padding:14px}
.kpi .k.warn{border-left-color:var(--warn)}
.kpi .lbl{font-size:11px;color:var(--muted);text-transform:uppercase;letter-spacing:.5px}
.kpi .val{font-size:20px;font-weight:700;margin-top:5px}
.kpi .hint{font-size:11px;color:var(--muted);margin-top:2px}
table{width:100%;border-collapse:collapse;font-size:13px}
th,td{text-align:left;padding:8px 9px;border-bottom:1px solid var(--line)}
th{font-size:11px;text-transform:uppercase;letter-spacing:.5px;color:var(--muted)}
td.num,th.num{text-align:right;font-variant-numeric:tabular-nums}
.bar{background:#eef4f4;height:6px;border-radius:3px;overflow:hidden;margin-top:4px}
.bar i{display:block;height:100%;background:var(--mid)}
.tools{display:flex;gap:8px;margin-bottom:10px;flex-wrap:wrap}
input,select{padding:7px 9px;border:1px solid var(--line);border-radius:7px;font-size:13px}
input{flex:1;min-width:160px}
.tag{font-size:10px;padding:2px 6px;border-radius:4px;background:#eaf1fb;color:#1e4b8b}
.tag.pindah{background:#eaf1fb;color:#1e4b8b}
.tag.masuk{background:#e6f4ec;color:#1e7b3c}
.scroll{max-height:520px;overflow:auto}
footer{color:var(--muted);font-size:12px;padding:0 0 30px;text-align:center}
@media(max-width:760px){.kpi{grid-template-columns:repeat(2,1fr)}}
</style></head><body>
<header><div class="wrap"><h1>Laporan Pengeluaran</h1>
<p class="sub" id="sub"></p></div></header>
<main class="wrap">
<div class="kpi" id="kpi"></div>
<div class="card"><h2>Per kategori</h2><div id="kat"></div></div>
<div class="card"><h2>Per bulan</h2><div id="bulan"></div></div>
<div class="card"><h2>10 pengeluaran terbesar</h2><div id="top"></div></div>
<div class="card"><h2>Semua transaksi</h2>
<div class="tools"><input id="cari" placeholder="Cari keterangan / kategori...">
<select id="filter"><option value="">Semua kategori</option></select>
<span id="hitung" style="font-size:12px;color:var(--muted);align-self:center"></span></div>
<div class="scroll" id="tabel"></div></div>
<footer id="foot"></footer>
</main>
<script>
const DATA = __DATA__;
const rp = n => "Rp " + Math.round(n).toLocaleString("id-ID");
const pct = n => (n*100).toFixed(1) + "%";
const esc = s => String(s).replace(/[&<>"]/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));

function kartu(lbl,val,hint,cls){
  return '<div class="k '+(cls||'')+'"><div class="lbl">'+lbl+'</div><div class="val">'+val+
    '</div><div class="hint">'+(hint||'')+'</div></div>';
}

function gambar(){
  const t = DATA.total, n = DATA.jumlah;
  document.getElementById("sub").textContent =
    DATA.periode.dari.split("-").reverse().join("/") + " s/d " +
    DATA.periode.sampai.split("-").reverse().join("/") + "  -  " + n + " transaksi";
  document.getElementById("kpi").innerHTML =
    kartu("Total keluar", rp((t.belanja||0)+(t.pindah||0)), n+" transaksi","") +
    kartu("Belanja", rp(t.belanja||0), pct((t.belanja||0)/((t.belanja||0)+(t.pindah||0)||1))+" dari total","") +
    kartu("Pindah dana", rp(t.pindah||0), "top up / tarik tunai","warn") +
    kartu("Uang masuk", rp(t.masuk||0), "tidak dihitung belanja","");

  const kat = Object.entries(DATA.per_kat).sort((a,b)=>b[1].v-a[1].v);
  const maks = kat.length ? kat[0][1].v : 1;
  const totalB = t.belanja || 1;
  let h = '<table><thead><tr><th>Kategori</th><th class="num">Jml</th>'+
    '<th class="num">Total</th><th class="num">Porsi</th></tr></thead><tbody>';
  kat.forEach(([nama,v])=>{
    h += '<tr><td>'+esc(nama)+'<div class="bar"><i style="width:'+(v.v/maks*100).toFixed(1)+
      '%"></i></div></td><td class="num">'+v.n+'</td><td class="num"><b>'+rp(v.v)+
      '</b></td><td class="num">'+pct(v.v/totalB)+'</td></tr>';
  });
  document.getElementById("kat").innerHTML = h+'</tbody></table>';

  const bl = Object.entries(DATA.per_bulan).sort((a,b)=>a[0]<b[0]?-1:1);
  const maksB = Math.max(...bl.map(x=>x[1]),1);
  h = '<table><thead><tr><th>Bulan</th><th class="num">Total</th></tr></thead><tbody>';
  bl.forEach(([ym,v])=>{
    h += '<tr><td>'+ym+'<div class="bar"><i style="width:'+(v/maksB*100).toFixed(1)+
      '%"></i></div></td><td class="num"><b>'+rp(v)+'</b></td></tr>';
  });
  document.getElementById("bulan").innerHTML = h+'</tbody></table>';

  const top = DATA.baris.filter(b=>b.jenis_uang==="belanja")
    .sort((a,b)=>b.nominal-a.nominal).slice(0,10);
  h = '<table><thead><tr><th>Tanggal</th><th>Keterangan</th><th>Kategori</th>'+
    '<th class="num">Nominal</th></tr></thead><tbody>';
  top.forEach(b=>{
    h += '<tr><td>'+b.tanggal.split("-").reverse().join("/")+'</td><td>'+esc(b.keterangan.slice(0,52))+
      '</td><td>'+esc(b.kategori)+'</td><td class="num">'+rp(b.nominal)+'</td></tr>';
  });
  document.getElementById("top").innerHTML = h+'</tbody></table>';

  const sel = document.getElementById("filter");
  kat.forEach(([nama])=>{
    const o=document.createElement("option"); o.value=nama; o.textContent=nama; sel.appendChild(o);
  });
  document.getElementById("foot").textContent =
    "Dibuat oleh SpendingToolkit - semua diproses di komputer ini, tidak ada data yang diupload.";
  tabel();
}

function tabel(){
  const cari = (document.getElementById("cari").value||"").toLowerCase();
  const filter = document.getElementById("filter").value;
  const rows = DATA.baris.filter(b=>{
    if(filter && b.kategori!==filter) return false;
    if(cari && !(b.keterangan+" "+b.kategori).toLowerCase().includes(cari)) return false;
    return true;
  });
  let h = '<table><thead><tr><th>Tanggal</th><th>Keterangan</th><th>Kategori</th>'+
    '<th>Jenis</th><th class="num">Nominal</th></tr></thead><tbody>';
  rows.forEach(b=>{
    h += '<tr><td>'+b.tanggal.split("-").reverse().join("/")+(b.waktu?'<br><span style="font-size:11px;color:#6b8484">'+b.waktu+'</span>':'')+
      '</td><td>'+esc(b.keterangan.slice(0,60))+'</td><td>'+esc(b.kategori)+
      '</td><td><span class="tag '+b.jenis_uang+'">'+b.jenis_uang+'</span></td>'+
      '<td class="num">'+rp(b.nominal)+'</td></tr>';
  });
  document.getElementById("tabel").innerHTML = h+'</tbody></table>';
  document.getElementById("hitung").textContent = rows.length + " dari " + DATA.baris.length;
}

document.getElementById("cari").oninput = tabel;
document.getElementById("filter").onchange = tabel;
gambar();
</script></body></html>
"""


def tulis_html(baris, r, path):
    total = r["total"]
    data = {
        "periode": {"dari": baris[0]["tanggal"], "sampai": baris[-1]["tanggal"]},
        "jumlah": len(baris),
        "total": {k: v for k, v in total.items()},
        "per_kat": r["per_kat"],
        "per_bulan": r["per_bulan"],
        "baris": [{k: b[k] for k in ("tanggal", "waktu", "keterangan",
                                     "kategori", "jenis_uang", "nominal")} for b in baris],
    }
    html = HTML.replace("__DATA__", json.dumps(data, ensure_ascii=False))
    with io.open(path, "w", encoding="utf-8") as fh:
        fh.write(html)
    return path


def main():
    ap = argparse.ArgumentParser(description="Buat laporan dari transaksi.csv")
    ap.add_argument("csv", help="file transaksi.csv")
    ap.add_argument("--kategori", default=KATEGORI_DEFAULT, help="file aturan kategori")
    ap.add_argument("--out", default="", help="folder hasil")
    args = ap.parse_args()

    if not os.path.exists(args.csv):
        print("Tidak ada: " + args.csv)
        return 1

    baris, aturan = proses(args.csv, args.kategori)
    if not baris:
        print("Tidak ada baris valid. Cek CSV-nya (kolom tanggal & nominal).")
        return 1

    r = ringkas(baris)
    folder = args.out or os.path.dirname(os.path.abspath(args.csv))
    if not os.path.isdir(folder):
        os.makedirs(folder)

    judul("HASIL")
    print("  Transaksi        : %d" % len(baris))
    print("  Periode          : %s s/d %s" % (baris[0]["tanggal"], baris[-1]["tanggal"]))
    print("  Belanja          : %s" % rupiah(r["total"].get("belanja", 0)))
    print("  Pindah dana      : %s" % rupiah(r["total"].get("pindah", 0)))
    print("  Uang masuk       : %s" % rupiah(r["total"].get("masuk", 0)))
    print("  Aturan kategori  : %d baris" % len(aturan))

    judul("PER KATEGORI")
    for kat, v in sorted(r["per_kat"].items(), key=lambda x: -x[1]["v"]):
        print("  %-24s %-12s %d tx" % (kat, rupiah(v["v"]), v["n"]))

    tak = [b for b in baris if b["kategori"] == "Lain-lain"]
    if tak:
        judul("BELUM ADA KATEGORI (%d transaksi)" % len(tak))
        print("  Tambahkan kata kunci di config/kategori.txt, lalu jalankan ulang:")
        for b in tak[:12]:
            print("    %s  %-10s %s" % (b["tanggal"], rupiah(b["nominal"]),
                                        b["keterangan"][:44]))

    html = tulis_html(baris, r, os.path.join(folder, "laporan.html"))
    xlsx = tulis_excel(baris, r, os.path.join(folder, "tracker.xlsx"))

    judul("FILE")
    print("  " + html)
    print("  " + xlsx)
    return 0


if __name__ == "__main__":
    sys.exit(main())
